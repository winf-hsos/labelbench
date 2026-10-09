"""Reading and writing the CSV, YAML and JSON files labelbench works with."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import sys
from pathlib import Path
from typing import Any

import yaml

# Text fields may be very long; the default limit of 128 KB is too small.
csv.field_size_limit(min(sys.maxsize, 2**31 - 1))


def read_csv(path: str | Path) -> list[dict[str, str]]:
    """Read a CSV file into a list of rows. A UTF-8 byte order mark is ignored."""
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def csv_columns(path: str | Path) -> list[str]:
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return next(csv.reader(fh), [])


TABLE_SUFFIXES = (".csv", ".xlsx")


def split_table_spec(spec: str | Path) -> tuple[Path, str | None]:
    """'data.xlsx#Sheet' -> (Path('data.xlsx'), 'Sheet'); without '#' the sheet is None."""
    text = str(spec)
    if "#" in text:
        path, sheet = text.rsplit("#", 1)
        return Path(path), sheet or None
    return Path(text), None


def table_file(spec: str | Path) -> Path:
    """The file behind a table spec, without any '#Sheet' suffix."""
    return split_table_spec(spec)[0]


def read_table(spec: str | Path) -> list[dict[str, str]]:
    """Read a CSV file or an Excel sheet into a list of rows with text values.

    The file extension decides: `.csv` is read as UTF-8 CSV, `.xlsx` with
    openpyxl. For Excel files `file.xlsx#Sheet` selects a sheet; without it
    the first sheet is used. The first row holds the column names.
    """
    path, sheet = split_table_spec(spec)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        if sheet:
            raise ValueError(f"{spec}: a CSV file has no sheets.")
        return read_csv(path)
    if suffix == ".xlsx":
        header, rows = _read_xlsx(path, sheet)
        return [dict(zip(header, r, strict=True)) for r in rows]
    raise ValueError(f"{path}: unsupported table format; use {' or '.join(TABLE_SUFFIXES)}.")


def table_columns(spec: str | Path) -> list[str]:
    path, sheet = split_table_spec(spec)
    if path.suffix.lower() == ".xlsx":
        return _read_xlsx(path, sheet, header_only=True)[0]
    return csv_columns(path)


def _read_xlsx(path: Path, sheet: str | None, header_only: bool = False) -> tuple[list[str], list[list[str]]]:
    import openpyxl  # imported here so CSV-only use does not need it

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet is not None and sheet not in wb.sheetnames:
            raise ValueError(f"{path} has no sheet {sheet!r}; sheets: {wb.sheetnames}")
        ws = wb[sheet] if sheet is not None else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        first = next(it, None) or ()
        # columns without a name are ignored, as are cells to their right that stay unnamed
        named = [(i, _cell_text(v).strip()) for i, v in enumerate(first) if _cell_text(v).strip()]
        header = [name for _, name in named]
        dup = sorted({h for h in header if header.count(h) > 1})
        if dup:
            raise ValueError(f"{path}#{ws.title}: duplicate column names {dup}")
        if header_only:
            return header, []
        rows = []
        for values in it:
            cells = [_cell_text(values[i]) if i < len(values) else "" for i, _ in named]
            if any(cells):
                rows.append(cells)
        return header, rows
    finally:
        wb.close()


def _cell_text(value: Any) -> str:
    """Excel cell as text: whole numbers without '.0', dates in ISO format, empty cells as ''."""
    import datetime as dt

    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ") if value.time() != dt.time() else value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def write_csv(path: str | Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    with Path(path).open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read_yaml(path: str | Path) -> dict[str, Any]:
    return yaml.safe_load(Path(path).read_text(encoding="utf-8")) or {}


def write_yaml(path: str | Path, data: dict[str, Any]) -> None:
    text = yaml.safe_dump(data, allow_unicode=True, sort_keys=False)
    Path(path).write_text(text, encoding="utf-8")


def read_json(path: str | Path) -> Any:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: str | Path, data: Any) -> None:
    text = json.dumps(data, indent=2, ensure_ascii=False)
    Path(path).write_text(text + "\n", encoding="utf-8")


def write_text_atomic(path: str | Path, text: str) -> None:
    """Write via a temporary file so an interrupted write never leaves half a file."""
    path = Path(path)
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    os.replace(tmp, path)


def sha256_file(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()
